import openpyxl
import json
import os

def parse_all_nursery_maps(excel_path, output_dir):
    if not os.path.exists(excel_path):
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    os.makedirs(output_dir, exist_ok=True)

    for cell_name in wb.sheetnames:
        if cell_name in ['Blad1', 'Blad2', 'Pesticides', 'Legenda plagen']: 
            continue
        
        ws = wb[cell_name]
        
        # Crete an empty grid 
        actual_max_r = min(ws.max_row, 79)
        actual_max_c = ws.max_column

        grid = []

        for r in range(actual_max_r):            
            new_row = [""] * actual_max_c
            grid.append(new_row)   

        # Fill basic data
        # Handle start of counting in Excel from 1 
        for r in range(1, actual_max_r + 1):
            for c in range(1, actual_max_c + 1):
                val = ws.cell(row=r, column=c).value # extract the data
                if val is not None:
                    if isinstance(val, (int, float)) and not isinstance(val, bool):
                        if float(val).is_integer():
                            grid[r-1][c-1] = int(val)
                        else:
                            grid[r-1][c-1] = val
                    else:
                        grid[r-1][c-1] = str(val) # the data is a string eg. 'kas'
        
        # Add information to the missing data points
        for merge_range in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merge_range.bounds
            # Get the top-left value in Excel since i tis the only cell with the value in Excel
            top_left_value = ws.cell(row=min_row, column=min_col).value
            
            fill_val = ""
            if top_left_value is not None:
                if isinstance(top_left_value, (int, float)):
                    if float(top_left_value).is_integer():
                        fill_val = int(top_left_value)
                    else:
                        fill_val = top_left_value
                else:
                    fill_val = str(top_left_value)
            
            # Overwrite with merged cell values
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    if r <= actual_max_r and c <= actual_max_c:
                        grid[r-1][c-1] = fill_val
        
        # Separately handle the title
        if len(grid) > 0:
            for c in range(len(grid[0])):
                grid[0][c] = "Legmeerdijk"
                

        # for r in range(22, 26):
        #     if r < len(grid) and len(grid[r]) > 1:
        #         grid[r][1] = "silo"

        # Save to the json file 
        file_path = os.path.join(output_dir, "leg.json")
        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(grid, f, indent=2)
        
if __name__ == "__main__":
    script_dir = os.path.dirname(os.path.abspath(__file__))
    excel_path = os.path.normpath(os.path.join(script_dir, '..', 'Maps', 'Legmeerdijk.xlsx'))
    output_dir = os.path.normpath(os.path.join(script_dir, '..', '..', 'frontend', 'src', 'components', 'AllMaps'))
    parse_all_nursery_maps(excel_path, output_dir)