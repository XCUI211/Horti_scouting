import openpyxl
import json
import os

def parse_all_nursery_maps(excel_path, output_dir, base_filename):
    if not os.path.exists(excel_path):
        print(f"Could not find {excel_path}")
        return

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    os.makedirs(output_dir, exist_ok=True)

    for sheet_name in wb.sheetnames:
        if sheet_name in ['Blad1', 'Blad2', 'Pesticides', 'Legenda plagen']: 
            continue
        
        ws = wb[sheet_name]
        temp_max_r = ws.max_row
        temp_max_c = ws.max_column

        def is_col_empty(col_idx, total_rows):
            for r in range(1, total_rows + 1):
                if ws.cell(row=r, column=col_idx).value is not None:
                    return False
            for merge_range in ws.merged_cells.ranges:
                min_col, _, max_col, _ = merge_range.bounds
                if min_col <= col_idx <= max_col:
                    return False
            return True

        actual_max_c = temp_max_c
        while actual_max_c > 0 and is_col_empty(actual_max_c, temp_max_r):
            actual_max_c -= 1

        actual_max_r = temp_max_r
        while actual_max_r > 0:
            row_has_data = any(ws.cell(row=actual_max_r, column=c).value is not None for c in range(1, actual_max_c + 1))
            row_has_merge = any(m.bounds[1] <= actual_max_r <= m.bounds[3] for m in ws.merged_cells.ranges)
            if row_has_data or row_has_merge:
                break
            actual_max_r -= 1

        if actual_max_r == 0 or actual_max_c == 0:
            continue

        grid = [[None for _ in range(actual_max_c)] for _ in range(actual_max_r)]
        
        for r in range(1, actual_max_r + 1):
            for c in range(1, actual_max_c + 1):
                grid[r-1][c-1] = ws.cell(row=r, column=c).value
        
        for merge_range in ws.merged_cells.ranges:
            min_col, min_row, max_col_range, max_row_range = merge_range.bounds
            top_left_value = ws.cell(row=min_row, column=min_col).value
            
            for r in range(min_row, max_row_range + 1):
                for c in range(min_col, max_col_range + 1):
                    if r <= actual_max_r and c <= actual_max_c:
                        grid[r-1][c-1] = top_left_value
        

        json_base = base_filename.replace(".xlsx", "").replace(" ", "_").lower()
        file_path = os.path.join(output_dir, f"{json_base}.json")
        
        with open(file_path, 'w', encoding='utf-8') as f:
            json.dump(grid, f, indent=2)
            


if __name__ == "__main__":
    script_dir = os.path.dirname(os.path.abspath(__file__))
    maps_dir = os.path.normpath(os.path.join(script_dir, '..', 'Maps'))
    output_dir = os.path.normpath(os.path.join(script_dir, '..', '..', 'frontend', 'src', 'components', 'AllMaps'))
    
    if not os.path.exists(maps_dir):
        print(f"Error: Can not find {maps_dir}")
    else:
        for file_name in os.listdir(maps_dir):
            if file_name.endswith('.xlsx') and not file_name.startswith('~$'):
                excel_path = os.path.join(maps_dir, file_name)
                parse_all_nursery_maps(excel_path, output_dir, file_name)

    print(f"\nAll maps are stored in: {output_dir}")
